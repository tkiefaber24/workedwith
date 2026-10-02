import io

MAX_RESUME_TEXT_CHARS = 50000


def _extract_pdf_text(file_bytes):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            pass
    return "\n".join(parts)


def _extract_docx_text(file_bytes):
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_resume_text(file_bytes, content_type):
    """Best-effort plain-text extraction for keyword search. Never raises --
    a resume that fails to parse just isn't searchable by its contents."""
    try:
        if content_type == "application/pdf":
            text = _extract_pdf_text(file_bytes)
        elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            text = _extract_docx_text(file_bytes)
        else:
            text = ""
    except Exception:
        text = ""
    return text[:MAX_RESUME_TEXT_CHARS]
